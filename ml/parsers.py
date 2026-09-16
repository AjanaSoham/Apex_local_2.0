from pathlib import Path

import pymupdf
try:
    import pytesseract
    from PIL import Image
except ImportError:
    pytesseract = None
    Image = None
from docx import Document


if pytesseract is not None:
    pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def extract_text_from_pdf(file_path: str) -> str:
    document = pymupdf.open(file_path)

    normal_text = ""

    for page in document:
        normal_text += page.get_text() + "\n"

    normal_text = normal_text.strip()

    # Use normal extraction if enough text is available.
    if len(normal_text) >= 50:
        document.close()
        return normal_text

    # Otherwise, use OCR when it is installed.
    if pytesseract is None or Image is None:
        document.close()
        raise ValueError("This scanned PDF needs OCR, but the optional OCR dependencies are unavailable.")
    print("Little or no text found. Running OCR...")

    ocr_text = []

    for page_number, page in enumerate(document):
        print(f"Running OCR on page {page_number + 1}...")

        pixmap = page.get_pixmap(
            matrix=pymupdf.Matrix(2, 2)
        )

        image = Image.frombytes(
            "RGB",
            [pixmap.width, pixmap.height],
            pixmap.samples
        )

        page_text = pytesseract.image_to_string(image)

        ocr_text.append(page_text)

    document.close()

    return "\n".join(ocr_text).strip()


def extract_text_from_docx(file_path: str) -> str:
    document = Document(file_path)

    extracted_text = []

    # Extract normal paragraphs.
    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            extracted_text.append(paragraph_text)

    # Extract tables.
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                extracted_text.append(" | ".join(row_text))

    return "\n".join(extracted_text).strip()


def extract_text(file_path: str) -> str:
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    elif extension == ".docx":
        return extract_text_from_docx(file_path)

    elif extension == ".txt":
        return path.read_text(encoding="utf-8", errors="replace").strip()

    else:
        raise ValueError(
            "Unsupported file type. "
            "Only PDF, DOCX and TXT files are supported."
        )
